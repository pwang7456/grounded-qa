"""docx 抽取：按文档体真实顺序产出 Block（标题层级来自样式名，表格转键值对行）。"""
import re

import blocks

HEADING_STYLE_RE = re.compile(r"(?:Heading|heading|标题)\s*(\d)")


def _para_block(para) -> dict:
    text = para.text.strip()
    if not text:
        return None
    style_name = ""
    try:
        style_name = str(para.style.name or "")
    except Exception:
        pass  # 损坏样式不应让整份文档入库失败
    m = HEADING_STYLE_RE.search(style_name)
    if m:
        return blocks.heading(text, int(m.group(1)))
    return blocks.paragraph(text)


def _table_block(table):
    if not table.rows:
        return None
    header = [c.text for c in table.rows[0].cells]
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    return blocks.table_from_rows(header, rows)


def extract_docx_blocks(path: str) -> list:
    """返回按原文顺序排列的 List[Block]。

    已知限制：文本框（Content Control / Drawing）与图片内的文字 python-docx 读不到；
    合并单元格会把同一文本重复读出，这里不做去重（真实语料遇到再处理）。
    """
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        raise ImportError(
            f"python-docx 未安装（{e}）；请先 python -m pip install -r requirements.txt")

    doc = Document(path)
    out = []
    # doc.paragraphs 与 doc.tables 是两个独立列表，直接用会丢掉段落与表格的交错顺序
    for child in doc.element.body.iterchildren():
        tag = child.tag
        if tag.endswith("}p"):
            blk = _para_block(Paragraph(child, doc))
        elif tag.endswith("}tbl"):
            blk = _table_block(Table(child, doc))
        else:
            continue
        if blk:
            out.append(blk)
    return out
